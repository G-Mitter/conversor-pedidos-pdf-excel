import glob
import os
import re
from pypdf import PdfReader
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import traceback

def processar_pdf(caminho_pdf):
    caminho_excel = os.path.splitext(caminho_pdf)[0] + ".xlsx"
    reader = PdfReader(caminho_pdf)
    
    branches = {}
    current_cnpj = None
    
    for page in reader.pages:
        text = page.extract_text()
        if not text:
            continue
            
        lines = [l.strip() for l in text.split('\n') if l.strip()]
        buyer_found = False # Pega APENAS o primeiro CNPJ da página (Comprador/Filial)
        
        for i, line in enumerate(lines):
            # 1. Captura a Filial (Comprador) e ignora o Fornecedor (2º bloco)
            if "CNPJ:" in line and "Insc. Estad." in line and not buyer_found:
                razao_social = line.split("CNPJ:")[0].strip()
                cnpj_match = re.search(r'CNPJ:\s*([\d\.\/-]+)', line)
                cnpj = cnpj_match.group(1) if cnpj_match else ""
                ie_match = re.search(r'Insc\. Estad\.:\s*([\d-]+)', line)
                ie = ie_match.group(1) if ie_match else ""
                
                endereco = ""
                telefone = ""
                if i + 1 < len(lines):
                    addr_line = lines[i+1]
                    if "CEP:" in addr_line or "Telefone:" in addr_line:
                        endereco = addr_line.split("CEP:")[0].strip() if "CEP:" in addr_line else addr_line.split("Telefone:")[0].strip()
                        tel_match = re.search(r'Telefone:\s*([\d\s()-]+)', addr_line)
                        telefone = tel_match.group(1).strip() if tel_match else ""
                        
                email = ""
                if i + 2 < len(lines):
                    email_line = lines[i+2]
                    if "Email:" in email_line:
                        email_match = re.search(r'Email:\s*([^\s]+)', email_line)
                        email = email_match.group(1).strip() if email_match else ""
                        
                buyer_found = True # Trava para não ler o CNPJ do fornecedor abaixo
                
                if cnpj not in branches:
                    branches[cnpj] = {
                        'razao_social': razao_social,
                        'cnpj': cnpj,
                        'ie': ie,
                        'endereco': endereco,
                        'telefone': telefone,
                        'email': email,
                        'produtos': []
                    }
                current_cnpj = cnpj
                
            # 2. Extração dos Produtos
            m_cod = re.match(r'^(\d{4,8})\s+(.*)', line)
            if m_cod and current_cnpj:
                cod_forn = m_cod.group(1)
                rest = m_cod.group(2)
                
                ean_match = re.search(r'(7\d{12})', rest)
                ean = ean_match.group(1) if ean_match else ''
                
                # Busca a UM exatamente no local da coluna, antes do multiplicador (ex: "CX  1")
                um_match = re.search(r'\b(CX|UN|PC|KG|FD|RL|M2|PAR|JG)\b(?=\s+\d+)', rest)
                if not um_match:
                    um_match = re.search(r'(CX|UN|PC|KG|FD|RL|M2|PAR|JG)(?=\s+\d+)', rest)
                    
                if um_match:
                    um = um_match.group(1)
                    desc_raw = rest[:um_match.start()].strip()
                    
                    # Remove o código do produto no ERP se estiver colado no final (ex: "123456", "AB1234")
                    desc = re.sub(r'\s+[A-Z0-9]+$', '', desc_raw).strip()
                    if not desc:
                        desc = desc_raw
                        
                    after_um = rest[um_match.end():].strip()
                    if ean:
                        after_um = after_um.replace(f"0,00{ean}", "0,00 ").replace(f"0,000{ean}", "0,000 ").replace(ean, "")
                        
                    tokens = [t for t in after_um.split() if re.match(r'^\d+(?:[\.,]\d+)*$', t)]
                    
                    quant, price, total = "", "", ""
                    
                    for idx_t, t in enumerate(tokens):
                        if re.match(r'^\d+,\d{3}$', t) and t != '0,000':
                            quant = t
                            if idx_t + 1 < len(tokens):
                                price = tokens[idx_t + 1]
                            break
                            
                    if not quant and len(tokens) >= 2:
                        quant = tokens[1]
                        if len(tokens) >= 3:
                            price = tokens[2]
                            
                    try:
                        q_v = float(quant.replace('.', '').replace(',', '.'))
                        p_v = float(price.replace('.', '').replace(',', '.'))
                        tot_v = q_v * p_v
                        total = f"{tot_v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
                    except:
                        pass
                        
                    if quant:
                        branches[current_cnpj]['produtos'].append({
                            'cod_forn': cod_forn,
                            'cod_barras': ean,
                            'descricao': desc,
                            'um': um,
                            'quant': quant,
                            'preco_unit': price,
                            'total': total
                        })

    # 3. Criar Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Pedidos por Filial"
    
    row_num = 1
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    bold_font = Font(bold=True)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    for cnpj, filial in branches.items():
        if not filial['produtos']:
            continue
            
        ws.cell(row=row_num, column=1, value="Nome da Filial:").font = bold_font
        ws.cell(row=row_num, column=2, value=filial['razao_social'])
        row_num += 1
        
        ws.cell(row=row_num, column=1, value="CNPJ:").font = bold_font
        ws.cell(row=row_num, column=2, value=filial['cnpj'])
        ws.cell(row=row_num, column=3, value="Inscrição Estadual:").font = bold_font
        ws.cell(row=row_num, column=4, value=filial['ie'])
        row_num += 1
        
        ws.cell(row=row_num, column=1, value="Endereço:").font = bold_font
        ws.cell(row=row_num, column=2, value=filial['endereco'])
        row_num += 1
        
        ws.cell(row=row_num, column=1, value="Telefone:").font = bold_font
        ws.cell(row=row_num, column=2, value=filial['telefone'])
        ws.cell(row=row_num, column=3, value="Email:").font = bold_font
        ws.cell(row=row_num, column=4, value=filial['email'])
        row_num += 2
        
        headers = ['Cód. Forn.', 'Código de Barras', 'Descrição', 'UM', 'Quant.', 'Preço Unit.', 'Total']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
        row_num += 1
        
        for prod in filial['produtos']:
            cells = [
                ws.cell(row=row_num, column=1, value=prod['cod_forn']),
                ws.cell(row=row_num, column=2, value=prod['cod_barras']),
                ws.cell(row=row_num, column=3, value=prod['descricao']),
                ws.cell(row=row_num, column=4, value=prod['um']),
                ws.cell(row=row_num, column=5, value=prod['quant']),
                ws.cell(row=row_num, column=6, value=prod['preco_unit']),
                ws.cell(row=row_num, column=7, value=prod['total'])
            ]
            for cell in cells:
                cell.border = border
            row_num += 1
            
        row_num += 2

    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 50)
        
    wb.save(caminho_excel)
    print(f"✅ Sucesso: '{caminho_pdf}' exportado para '{caminho_excel}' ({len(branches)} filiais encontradas)")

def processar_todos_os_pdfs():
    arquivos_pdf = glob.glob("*.pdf")
    
    if not arquivos_pdf:
        print("Nenhum arquivo .pdf foi encontrado na pasta atual.")
        return
        
    print(f"Foram encontrados {len(arquivos_pdf)} arquivo(s) PDF para processar...")
    for pdf in arquivos_pdf:
        try:
            processar_pdf(pdf)
        except Exception as e:
            print(f"❌ Erro ao processar o arquivo '{pdf}': {e}")
            traceback.print_exc()

if __name__ == "__main__":
    processar_todos_os_pdfs()